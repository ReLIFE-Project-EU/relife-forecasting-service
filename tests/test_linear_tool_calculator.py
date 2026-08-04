from __future__ import annotations

from datetime import datetime
from pathlib import Path

import openpyxl
import pytest
from fastapi.testclient import TestClient

from relife_forecasting.main import app
from relife_forecasting.scripts.linear_tool_calculator import (
    calculate,
    calculate_linear_heat_cold_daly,
    load_scenarios,
    read_temperatures,
    read_temperatures_csv,
)


WORKBOOK_PATH = Path("D_Antolucci_NEI_CoBen_linear_heat_cold_DALY_tool_v2.xlsx")
client = TestClient(app)

requires_workbook = pytest.mark.skipif(
    not WORKBOOK_PATH.exists(),
    reason="Source workbook is not committed to the repository.",
)


@requires_workbook
def test_linear_tool_calculator_matches_workbook_formula_outputs():
    wb = openpyxl.load_workbook(WORKBOOK_PATH, data_only=True)
    scenarios = load_scenarios(wb["Parameters"])
    result = calculate(read_temperatures(wb["Linear_tool"]), scenarios["COMFORT_PAIR_26_20"])

    assert result["selected_threshold_pair"] == "COMFORT_PAIR_26_20"
    assert result["scenario_label"] == "Indoor comfort pair"
    assert result["valid_daily_temperature_rows"] == 279
    assert result["summary_average_days"] == 279
    assert result["average_daily_heat_harm_rate"] == 1.0615072320803287
    assert result["average_daily_cold_harm_rate"] == 0.0
    assert result["average_daily_total_harm_rate"] == 1.0615072320803287
    assert result["annual_period_total_harm_for_population"] == 0.0038214260354891834
    assert result["rows"][0]["thermal_atribution"] == "Zero"
    assert result["rows"][-1]["date"] == datetime(2026, 10, 6)


@requires_workbook
def test_linear_tool_calculator_can_reproduce_saved_excel_summary_divisor():
    wb = openpyxl.load_workbook(WORKBOOK_PATH, data_only=True)
    scenarios = load_scenarios(wb["Parameters"])
    result = calculate(
        read_temperatures(wb["Linear_tool"]),
        scenarios["COMFORT_PAIR_26_20"],
        average_days=95,
    )

    assert result["valid_daily_temperature_rows"] == 279
    assert result["summary_average_days"] == 95
    assert result["average_daily_heat_harm_rate"] == pytest.approx(3.11747913421486)
    assert result["average_daily_total_harm_per_person"] == pytest.approx(0.0000311747913421486)
    assert result["pasted_profile_heat_harm_for_population"] == pytest.approx(0.002961605177504116)
    assert result["annual_period_total_harm_for_population"] == pytest.approx(0.011222924883173495)


@requires_workbook
def test_linear_tool_calculator_accepts_csv_input(tmp_path):
    csv_path = tmp_path / "temps.csv"
    csv_path.write_text(
        "date,indoor_air_temperature_C\n"
        "2026-01-01,22.0\n"
        "2026-01-02,27.0\n",
        encoding="utf-8",
    )

    wb = openpyxl.load_workbook(WORKBOOK_PATH, data_only=True)
    scenarios = load_scenarios(wb["Parameters"])
    temperatures = read_temperatures_csv(csv_path)
    result = calculate(temperatures, scenarios["COMFORT_PAIR_26_20"])

    assert [row["thermal_atribution"] for row in result["rows"]] == ["Zero", "Heat"]
    assert result["average_daily_heat_harm_rate"] == (0.0 + (27.0 - 26.0) * 0.7152769) / 2


def test_linear_tool_calculator_accepts_csv_decimal_comma(tmp_path):
    csv_path = tmp_path / "temps_decimal_comma.csv"
    csv_path.write_text(
        "date,indoor_air_temperature_C\n"
        "2026-01-01,24,29\n",
        encoding="utf-8",
    )

    assert read_temperatures_csv(csv_path) == [(datetime(2026, 1, 1), 24.29)]


def test_calculate_linear_heat_cold_daly_uses_excel_scenario_options():
    result = calculate_linear_heat_cold_daly(
        [22.0, 27.0],
        selected_threshold_pair="COMFORT_PAIR_26_20",
        scenario_label="Indoor comfort pair",
        population_persons=1,
        exposure_days_for_period=360,
        conversion_factor_persons=100000,
        heat_threshold_c="26,0",
        cold_threshold_c="20,0",
        heat_linear_hi="0,7152769",
        cold_linear_hi="0,3481684",
        neutral_zone_note="Zero harm between 20.0 °C and 26.0 °C.",
    )

    assert result["selected_threshold_pair"] == "COMFORT_PAIR_26_20"
    assert result["scenario_label"] == "Indoor comfort pair"
    assert result["heat_threshold_C"] == 26.0
    assert result["cold_threshold_C"] == 20.0
    assert result["average_daily_heat_harm_rate"] == pytest.approx(0.35763845)


def test_calculate_linear_heat_cold_daly_rejects_mixed_scenario_values():
    with pytest.raises(ValueError, match="heat_threshold_C"):
        calculate_linear_heat_cold_daly(
            [27.0],
            selected_threshold_pair="COMFORT_PAIR_26_20",
            heat_threshold_c=30.0,
        )


def test_calculate_linear_heat_cold_daly_rejects_unknown_scenario():
    with pytest.raises(ValueError, match="Unknown selected_threshold_pair"):
        calculate_linear_heat_cold_daly([27.0], selected_threshold_pair="BAD_SCENARIO")  # type: ignore[arg-type]


def test_linear_tool_api_endpoint_matches_calculator():
    payload = {
        "temperatures_c": [
            ["2026-01-01T00:00:00", 22.0],
            ["2026-01-02T00:00:00", 27.0],
        ],
        "selected_threshold_pair": "COMFORT_PAIR_26_20",
        "scenario_label": "Indoor comfort pair",
        "population_persons": 1,
        "exposure_days_for_period": 360,
        "conversion_factor_persons": 100000,
        "heat_threshold_c": "26,0",
        "cold_threshold_c": "20,0",
        "heat_linear_hi": "0,7152769",
        "cold_linear_hi": "0,3481684",
        "neutral_zone_note": "Zero harm between 20.0 °C and 26.0 °C.",
    }

    response = client.post("/linear-tool/heat-cold-daly", json=payload)

    assert response.status_code == 200
    data = response.json()
    assert data["selected_threshold_pair"] == "COMFORT_PAIR_26_20"
    assert data["scenario_label"] == "Indoor comfort pair"
    assert data["heat_threshold_C"] == 26.0
    assert data["cold_threshold_C"] == 20.0
    assert data["average_daily_heat_harm_rate"] == pytest.approx(0.35763845)
