#!/usr/bin/env python3
"""Replicate the Excel `Linear_tool` sheet in Python.

The workbook contains a fixed scenario table and a daily temperature profile.
This script reads the workbook, applies the same row-wise formulas used in the
sheet, and prints a compact summary. It can also export the per-day results.
"""

from __future__ import annotations

import argparse
import csv
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Iterable, Literal

import openpyxl


@dataclass(frozen=True)
class Scenario:
    scenario_id: str
    scenario_label: str
    heat_threshold_c: float
    cold_threshold_c: float
    heat_linear_hi: float
    cold_linear_hi: float


ThresholdPair = Literal["ITALY_MMT_24_4", "COMFORT_PAIR_26_20", "CVD_PAIR_30_21"]


EXCEL_SCENARIOS: dict[ThresholdPair, Scenario] = {
    "ITALY_MMT_24_4": Scenario(
        scenario_id="ITALY_MMT_24_4",
        scenario_label="Italy MMT pair",
        heat_threshold_c=24.4,
        cold_threshold_c=24.4,
        heat_linear_hi=0.7152769,
        cold_linear_hi=0.3481684,
    ),
    "COMFORT_PAIR_26_20": Scenario(
        scenario_id="COMFORT_PAIR_26_20",
        scenario_label="Indoor comfort pair",
        heat_threshold_c=26.0,
        cold_threshold_c=20.0,
        heat_linear_hi=0.7152769,
        cold_linear_hi=0.3481684,
    ),
    "CVD_PAIR_30_21": Scenario(
        scenario_id="CVD_PAIR_30_21",
        scenario_label="CVD / health threshold pair",
        heat_threshold_c=30.0,
        cold_threshold_c=21.0,
        heat_linear_hi=0.7152769,
        cold_linear_hi=0.3481684,
    ),
}


def load_scenarios(ws) -> dict[str, Scenario]:
    scenarios: dict[str, Scenario] = {}
    for row in range(4, 7):
        scenario = Scenario(
            scenario_id=str(ws[f"A{row}"].value),
            scenario_label=str(ws[f"B{row}"].value),
            heat_threshold_c=float(ws[f"C{row}"].value),
            cold_threshold_c=float(ws[f"D{row}"].value),
            heat_linear_hi=float(ws[f"E{row}"].value),
            cold_linear_hi=float(ws[f"F{row}"].value),
        )
        scenarios[scenario.scenario_id] = scenario
    return scenarios


def _assert_matches(name: str, provided: str | float | None, expected: str | float) -> None:
    if provided is None:
        return
    if isinstance(expected, float):
        if abs(parse_float(provided) - expected) > 1e-9:
            raise ValueError(f"{name}={provided!r} does not match selected scenario value {expected!r}")
        return
    if str(provided) != expected:
        raise ValueError(f"{name}={provided!r} does not match selected scenario value {expected!r}")


def _normalise_temperatures(
    temperatures: Iterable[float | int | str | tuple[datetime, float | int | str]],
) -> list[tuple[datetime, float]]:
    rows: list[tuple[datetime, float]] = []
    for idx, item in enumerate(temperatures, start=1):
        if isinstance(item, tuple):
            date_value, temp_value = item
            rows.append((date_value, parse_float(temp_value)))
        else:
            rows.append((datetime(1900, 1, 1), parse_float(item)))
    return rows


def calculate_linear_heat_cold_daly(
    temperatures_c: Iterable[float | int | str | tuple[datetime, float | int | str]],
    selected_threshold_pair: ThresholdPair = "COMFORT_PAIR_26_20",
    population_persons: float = 1.0,
    exposure_days_for_period: float = 360.0,
    conversion_factor_persons: float = 100000.0,
    valid_daily_temperature_rows: int | None = None,
    scenario_label: str | None = None,
    average_indoor_temperature_c: float | str | None = None,
    heat_threshold_c: float | str | None = None,
    cold_threshold_c: float | str | None = None,
    heat_linear_hi: float | str | None = None,
    cold_linear_hi: float | str | None = None,
    neutral_zone_note: str | None = None,
) -> dict:
    """Single public function for the Excel Linear_tool calculation.

    Valid selected_threshold_pair options:
    - "ITALY_MMT_24_4": Italy MMT pair, heat_threshold_C=24.4,
      cold_threshold_C=24.4, heat_linear_HI=0.7152769,
      cold_linear_HI=0.3481684.
    - "COMFORT_PAIR_26_20": Indoor comfort pair, heat_threshold_C=26.0,
      cold_threshold_C=20.0, heat_linear_HI=0.7152769,
      cold_linear_HI=0.3481684.
    - "CVD_PAIR_30_21": CVD / health threshold pair,
      heat_threshold_C=30.0, cold_threshold_C=21.0,
      heat_linear_HI=0.7152769, cold_linear_HI=0.3481684.

    Values such as labels, thresholds and HI coefficients are fixed by
    selected_threshold_pair. Optional matching arguments are accepted only as
    validation constraints, so callers cannot silently mix scenarios.
    """

    if selected_threshold_pair not in EXCEL_SCENARIOS:
        available = ", ".join(EXCEL_SCENARIOS)
        raise ValueError(f"Unknown selected_threshold_pair {selected_threshold_pair!r}. Available: {available}")

    scenario = EXCEL_SCENARIOS[selected_threshold_pair]
    neutral_note = (
        "No neutral band: heat and cold deviations use the same benchmark."
        if scenario.heat_threshold_c == scenario.cold_threshold_c
        else f"Zero harm between {scenario.cold_threshold_c:.1f} °C and {scenario.heat_threshold_c:.1f} °C."
    )

    _assert_matches("scenario_label", scenario_label, scenario.scenario_label)
    _assert_matches("heat_threshold_C", heat_threshold_c, scenario.heat_threshold_c)
    _assert_matches("cold_threshold_C", cold_threshold_c, scenario.cold_threshold_c)
    _assert_matches("heat_linear_HI", heat_linear_hi, scenario.heat_linear_hi)
    _assert_matches("cold_linear_HI", cold_linear_hi, scenario.cold_linear_hi)
    _assert_matches("neutral_zone_note", neutral_zone_note, neutral_note)

    rows = _normalise_temperatures(temperatures_c)
    result = calculate(
        rows,
        scenario,
        population_persons=population_persons,
        conversion_factor_persons=conversion_factor_persons,
        exposure_days_for_period=exposure_days_for_period,
        average_days=valid_daily_temperature_rows,
    )

    if average_indoor_temperature_c is not None:
        _assert_matches(
            "average_indoor_temperature_C",
            average_indoor_temperature_c,
            result["average_indoor_temperature_C"],
        )

    return result


def read_temperatures(ws) -> list[tuple[datetime, float]]:
    rows: list[tuple[datetime, float]] = []
    for row in range(27, ws.max_row + 1):
        date_value = ws[f"B{row}"].value
        temp_value = ws[f"C{row}"].value
        if date_value in (None, "") or temp_value in (None, ""):
            continue
        if not isinstance(date_value, datetime):
            raise TypeError(f"Expected datetime in B{row}, got {type(date_value)!r}")
        rows.append((date_value, float(temp_value)))
    return rows


def parse_float(value: str | int | float) -> float:
    if isinstance(value, int | float):
        return float(value)
    return float(str(value).strip().replace(",", "."))


def read_temperatures_csv(path: Path) -> list[tuple[datetime, float]]:
    rows: list[tuple[datetime, float]] = []
    with path.open("r", newline="", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        if not reader.fieldnames:
            raise ValueError("CSV is missing a header row")
        for line_no, row in enumerate(reader, start=2):
            normalized_row = {}
            extra_values = row.get(None) or []
            for key, value in row.items():
                if key is None:
                    continue
                normalized_row[key.lstrip("\ufeff")] = value
            date_raw = (
                normalized_row.get("date")
                or normalized_row.get("B")
                or normalized_row.get("day")
                or ""
            ).strip()
            temp_raw = (
                normalized_row.get("indoor_air_temperature_C")
                or normalized_row.get("temperature")
                or normalized_row.get("temp")
                or normalized_row.get("C")
                or ""
            ).strip()
            if extra_values and temp_raw:
                temp_raw = ",".join([temp_raw, *extra_values])
            if not date_raw and not temp_raw:
                continue
            if not date_raw or not temp_raw:
                raise ValueError(f"Missing date or temperature in CSV row {line_no}")
            date_value = datetime.fromisoformat(date_raw)
            rows.append((date_value, parse_float(temp_raw)))
    return rows


def calculate(
    temperatures: Iterable[tuple[datetime, float]],
    scenario: Scenario,
    population_persons: float = 1.0,
    conversion_factor_persons: float = 100000.0,
    exposure_days_for_period: float = 360.0,
    average_days: float | None = None,
) -> dict:
    rows = []
    temps = list(temperatures)
    valid_rows = len(temps)
    summary_days = average_days if average_days is not None else valid_rows
    avg_temp = sum(t for _, t in temps) / valid_rows if valid_rows else None

    daily_heat_sum = 0.0
    daily_cold_sum = 0.0
    daily_total_sum = 0.0
    for idx, (date_value, temp) in enumerate(temps, start=1):
        heat_deviation = max(0.0, temp - scenario.heat_threshold_c)
        cold_deviation = max(0.0, scenario.cold_threshold_c - temp)
        heat_harm_rate = heat_deviation * scenario.heat_linear_hi
        cold_harm_rate = cold_deviation * scenario.cold_linear_hi
        total_harm_rate = heat_harm_rate + cold_harm_rate
        heat_per_person = heat_harm_rate / conversion_factor_persons
        cold_per_person = cold_harm_rate / conversion_factor_persons
        total_per_person = total_harm_rate / conversion_factor_persons
        heat_pop_day = heat_harm_rate * population_persons / conversion_factor_persons
        cold_pop_day = cold_harm_rate * population_persons / conversion_factor_persons
        total_pop_day = total_harm_rate * population_persons / conversion_factor_persons
        if temp > scenario.heat_threshold_c:
            attribution = "Heat"
        elif temp < scenario.cold_threshold_c:
            attribution = "Cold"
        else:
            attribution = "Zero"

        daily_heat_sum += heat_harm_rate
        daily_cold_sum += cold_harm_rate
        daily_total_sum += total_harm_rate

        rows.append(
            {
                "day_id": idx,
                "date": date_value,
                "indoor_air_temperature_C": temp,
                "heat_threshold_C": scenario.heat_threshold_c,
                "cold_threshold_C": scenario.cold_threshold_c,
                "heat_deviation_C": heat_deviation,
                "cold_deviation_C": cold_deviation,
                "heat_HI": scenario.heat_linear_hi,
                "cold_HI": scenario.cold_linear_hi,
                "heat_harm_rate": heat_harm_rate,
                "cold_harm_rate": cold_harm_rate,
                "total_harm_rate": total_harm_rate,
                "heat_per_person": heat_per_person,
                "cold_per_person": cold_per_person,
                "total_per_person": total_per_person,
                "heat_population_DALY_day": heat_pop_day,
                "cold_population_DALY_day": cold_pop_day,
                "total_population_DALY_day": total_pop_day,
                "thermal_atribution": attribution,
            }
        )

    average_daily_heat_harm_rate = daily_heat_sum / summary_days if summary_days else None
    average_daily_cold_harm_rate = daily_cold_sum / summary_days if summary_days else None
    average_daily_total_harm_rate = daily_total_sum / summary_days if summary_days else None

    return {
        "selected_threshold_pair": scenario.scenario_id,
        "scenario_label": scenario.scenario_label,
        "population_persons": population_persons,
        "exposure_days_for_period": exposure_days_for_period,
        "conversion_factor_persons": conversion_factor_persons,
        "valid_daily_temperature_rows": valid_rows,
        "summary_average_days": summary_days,
        "average_indoor_temperature_C": avg_temp,
        "heat_threshold_C": scenario.heat_threshold_c,
        "cold_threshold_C": scenario.cold_threshold_c,
        "heat_linear_HI": scenario.heat_linear_hi,
        "cold_linear_HI": scenario.cold_linear_hi,
        "neutral_zone_note": (
            "No neutral band: heat and cold deviations use the same benchmark."
            if scenario.heat_threshold_c == scenario.cold_threshold_c
            else f"Zero harm between {scenario.cold_threshold_c:.1f} °C and {scenario.heat_threshold_c:.1f} °C."
        ),
        "average_daily_heat_harm_rate": average_daily_heat_harm_rate,
        "average_daily_cold_harm_rate": average_daily_cold_harm_rate,
        "average_daily_total_harm_rate": average_daily_total_harm_rate,
        "average_daily_heat_harm_per_person": (
            average_daily_heat_harm_rate / conversion_factor_persons
            if average_daily_heat_harm_rate is not None
            else None
        ),
        "average_daily_cold_harm_per_person": (
            average_daily_cold_harm_rate / conversion_factor_persons
            if average_daily_cold_harm_rate is not None
            else None
        ),
        "average_daily_total_harm_per_person": (
            average_daily_total_harm_rate / conversion_factor_persons
            if average_daily_total_harm_rate is not None
            else None
        ),
        "average_daily_heat_harm_for_population": (
            average_daily_heat_harm_rate * population_persons / conversion_factor_persons
            if average_daily_heat_harm_rate is not None
            else None
        ),
        "average_daily_cold_harm_for_population": (
            average_daily_cold_harm_rate * population_persons / conversion_factor_persons
            if average_daily_cold_harm_rate is not None
            else None
        ),
        "average_daily_total_harm_for_population": (
            average_daily_total_harm_rate * population_persons / conversion_factor_persons
            if average_daily_total_harm_rate is not None
            else None
        ),
        "annual_period_heat_harm_for_population": (
            (average_daily_heat_harm_rate * population_persons / conversion_factor_persons)
            * exposure_days_for_period
            if average_daily_heat_harm_rate is not None
            else None
        ),
        "annual_period_cold_harm_for_population": (
            (average_daily_cold_harm_rate * population_persons / conversion_factor_persons)
            * exposure_days_for_period
            if average_daily_cold_harm_rate is not None
            else None
        ),
        "annual_period_total_harm_for_population": (
            (average_daily_total_harm_rate * population_persons / conversion_factor_persons)
            * exposure_days_for_period
            if average_daily_total_harm_rate is not None
            else None
        ),
        "pasted_profile_heat_harm_for_population": sum(r["heat_population_DALY_day"] for r in rows),
        "pasted_profile_cold_harm_for_population": sum(r["cold_population_DALY_day"] for r in rows),
        "pasted_profile_total_harm_for_population": sum(r["total_population_DALY_day"] for r in rows),
        "rows": rows,
    }


def call_main_api_linear_heat_cold_daly(
    temperatures_c: Iterable[float | int | str | tuple[datetime, float | int | str]],
    *,
    base_url: str = "http://127.0.0.1:8000",
    endpoint_path: str = "/linear-tool/heat-cold-daly",
    selected_threshold_pair: ThresholdPair = "COMFORT_PAIR_26_20",
    population_persons: float = 1.0,
    exposure_days_for_period: float = 360.0,
    conversion_factor_persons: float = 100000.0,
    valid_daily_temperature_rows: int | None = None,
    scenario_label: str | None = None,
    average_indoor_temperature_c: float | str | None = None,
    heat_threshold_c: float | str | None = None,
    cold_threshold_c: float | str | None = None,
    heat_linear_hi: float | str | None = None,
    cold_linear_hi: float | str | None = None,
    neutral_zone_note: str | None = None,
    timeout_seconds: float = 30.0,
) -> dict:
    """
    Call the FastAPI endpoint exposed by main.py and return its JSON response.

    Useful for a quick smoke test that the API wiring matches the standalone calculator.
    """
    try:
        import requests
    except Exception as exc:  # pragma: no cover - import failure is environment-specific
        raise RuntimeError("The 'requests' package is required to call the API endpoint.") from exc

    payload = {
        "selected_threshold_pair": selected_threshold_pair,
        "population_persons": population_persons,
        "exposure_days_for_period": exposure_days_for_period,
        "conversion_factor_persons": conversion_factor_persons,
        "valid_daily_temperature_rows": valid_daily_temperature_rows,
        "scenario_label": scenario_label,
        "average_indoor_temperature_c": average_indoor_temperature_c,
        "heat_threshold_c": heat_threshold_c,
        "cold_threshold_c": cold_threshold_c,
        "heat_linear_hi": heat_linear_hi,
        "cold_linear_hi": cold_linear_hi,
        "neutral_zone_note": neutral_zone_note,
    }
    payload["temperatures_c"] = []
    for temp_row in temperatures_c:
        if isinstance(temp_row, tuple):
            date_value, temp_value = temp_row
            payload["temperatures_c"].append([date_value.isoformat(), temp_value])
        else:
            payload["temperatures_c"].append(temp_row)

    response = requests.post(
        f"{base_url.rstrip('/')}{endpoint_path}",
        json=payload,
        timeout=timeout_seconds,
    )
    response.raise_for_status()
    return response.json()


def write_csv(rows: list[dict], path: Path) -> None:
    if not rows:
        return
    with path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=list(rows[0].keys()))
        writer.writeheader()
        writer.writerows(rows)


def main() -> int:
    parser = argparse.ArgumentParser(description="Replicate the Excel Linear_tool sheet.")
    parser.add_argument(
        "--workbook",
        default="D_Antolucci_NEI_CoBen_linear_heat_cold_DALY_tool_v2.xlsx",
        help="Path to the source workbook.",
    )
    parser.add_argument(
        "--input-csv",
        default=None,
        help="Optional CSV with columns like date and indoor_air_temperature_C.",
    )
    parser.add_argument(
        "--scenario",
        default="COMFORT_PAIR_26_20",
        choices=list(EXCEL_SCENARIOS),
        help="Scenario id from the Parameters sheet.",
    )
    parser.add_argument(
        "--csv",
        default=None,
        help="Optional output CSV for the row-wise calculation.",
    )
    parser.add_argument(
        "--average-days",
        type=float,
        default=None,
        help="Optional divisor for average daily summary outputs. Use 95 to reproduce the saved Excel summary.",
    )
    args = parser.parse_args()

    if args.input_csv:
        temperatures = read_temperatures_csv(Path(args.input_csv))
    else:
        wb = openpyxl.load_workbook(args.workbook, data_only=True)
        temperatures = read_temperatures(wb["Linear_tool"])

    result = calculate(temperatures, EXCEL_SCENARIOS[args.scenario], average_days=args.average_days)

    print(f"Scenario: {result['selected_threshold_pair']} ({result['scenario_label']})")
    print(f"Rows: {result['valid_daily_temperature_rows']}")
    print(f"Summary average days: {result['summary_average_days']}")
    print(f"Average daily heat harm rate: {result['average_daily_heat_harm_rate']}")
    print(f"Average daily cold harm rate: {result['average_daily_cold_harm_rate']}")
    print(f"Average daily total harm rate: {result['average_daily_total_harm_rate']}")
    print(f"Pasted-profile total harm for population: {result['pasted_profile_total_harm_for_population']}")
    print(f"Annual/period total harm for population: {result['annual_period_total_harm_for_population']}")

    if args.csv:
        write_csv(result["rows"], Path(args.csv))
        print(f"Wrote {args.csv}")

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
